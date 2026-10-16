# -*- coding: utf-8 -*-
"""Отчёт КЦ из Google-таблицы (заполняется руками) → строки для сверки с Битриксом.

Таблица открыта по ссылке, скачивается как xlsx без авторизации (export?format=xlsx).
Лист = менеджер; колонки: Дата, №, ФИО, Телефон, Филиал, Статус, [Дата записи], Комментарий, Предоплата.
Ловушки, уже пойманные: телефон бывает числом «87015580881.0» (хвост .0 ломает последние 10 цифр);
листы «Инфо…», «Итог», «ArthroExpert» — не менеджеры; имена листов сокращённые («АйданаТ»),
буквы казахского алфавита (Шынғыс/Шынгыс) — сопоставляем с картой Битрикса по нормализованному имени.
Только чтение. Персональные данные наружу не отдаём — телефон маскируется на уровне ответа панели."""
import os, io, re, time, datetime
import urllib.request
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE_DIR = os.environ.get("KC_CACHE_DIR") or os.path.join(HERE, ".cache")     # …/Сверка/.cache
DEFAULT_URL = "https://docs.google.com/spreadsheets/d/1Kt0WwVHOhqwKFzO88PGlr1bep9wJ-fiAlGyTq7KBEI8/export?format=xlsx"
SKIP = ("Итог", "ArthroExpert")
_MEM = {"mtime": None, "data": None}
KZ = str.maketrans("ғқңәөұүіһҒҚҢӘӨҰҮІҺ", "гкнаоууихГКНАОУУИХ")


def norm_phone(p):
    s = str(p or "").strip()
    if s.endswith(".0") and s[:-2].replace(" ", "").isdigit():
        s = s[:-2]
    dg = re.sub(r"\D", "", s)
    if len(dg) >= 11 and dg.startswith("8"):
        dg = "7" + dg[1:]
    return dg[-10:] if len(dg) >= 10 else ""


def to_day(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    m = re.match(r"(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2,4}))?", str(v or "").strip())
    if not m:
        return ""
    d, mo = int(m.group(1)), int(m.group(2))
    y = int(m.group(3)) if m.group(3) else datetime.date.today().year
    if y < 100:
        y += 2000
    try:
        return datetime.date(y, mo, d).isoformat()
    except Exception:
        return ""


def nname(s):
    return re.sub(r"[^a-zа-яё]", "", str(s or "").lower().translate(KZ))


def map_sheet_to_user(sheet, users):
    """Имя листа → полное имя из карты Битрикса. 'АйданаТ' → 'Айдана Тажибай', 'Шынгыс' → 'Шынғыс Ерболатов'.
    Неоднозначно — None (лучше честное «не сопоставлен», чем чужие цифры)."""
    s = nname(sheet)
    full = {u: nname(u) for u in users}
    for u, n in full.items():
        if n == s:
            return u
    # «АйданаТ» / «АйданаЖунисбаева»: имя и хвост фамилии слиты — режем по заглавной букве ДО нормализации
    raw = str(sheet or "").strip()
    m = re.match(r"^([А-ЯЁӘҒҚҢӨҰҮІҺ][а-яёәғқңөұүіһa-z]+)(.*)$", raw)
    if m:
        first, rest = nname(m.group(1)), nname(m.group(2))
    else:
        first, rest = s, ""
    cands = []
    for u in users:
        parts = [nname(p) for p in str(u).split()]
        if not parts or parts[0] != first:
            continue
        if rest and not any(p.startswith(rest) for p in parts[1:]):
            continue
        cands.append(u)
    return cands[0] if len(cands) == 1 else None


def download(url=None, ttl=1800):
    url = url or DEFAULT_URL
    os.makedirs(CACHE_DIR, exist_ok=True)
    p = os.path.join(CACHE_DIR, "kc_sheet.xlsx")
    if os.path.exists(p) and time.time() - os.path.getmtime(p) < ttl:
        return p
    # Google отдаёт файл потоком и на нём регулярно рвётся соединение (IncompleteRead на 2 МБ,
    # 25.09.2026): без повторов сверка с отчётом КЦ молча выпадала из дневного отчёта.
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    data, last = b"", None
    for attempt in range(4):
        try:
            with urllib.request.urlopen(req, timeout=180) as r:
                data = r.read()
            break
        except Exception as e:
            last = e
            data = b""
            time.sleep(3 * (attempt + 1))
    if not data:
        raise RuntimeError("не удалось скачать таблицу КЦ за 4 попытки: %s" % repr(last)[:150])
    if not data.startswith(b"PK"):
        raise RuntimeError("Google вернул не xlsx — таблица закрыта по ссылке?")
    tmp = p + ".part"
    open(tmp, "wb").write(data)
    os.replace(tmp, p)
    return p


def load_rows(url=None, ttl=1800):
    """→ {"taken": 'дата-время скачивания', "rows": [...], "sheets": [...]}. Кэш в памяти по mtime файла."""
    p = download(url, ttl)
    mt = os.path.getmtime(p)
    if _MEM["mtime"] == mt and _MEM["data"]:
        return _MEM["data"]
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    rows, sheets = [], []
    for sn in wb.sheetnames:
        if sn in SKIP or sn.startswith("Инфо"):
            continue
        sheets.append(sn)
        ws = wb[sn]
        idx = {}
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                idx = {str(x or "").strip().lower(): j for j, x in enumerate(r)}
                continue
            if not r or not any(r):
                continue
            g = lambda k: (r[idx[k]] if k in idx and idx[k] < len(r) else None)
            # у части листов колонка телефона подписана «Номер» (лист «Сания Артро», 14.10.2026)
            if "телефон" not in idx:
                for alt in ("номер", "номер телефона", "тел", "тел."):
                    if alt in idx:
                        idx["телефон"] = idx[alt]; break
            day = to_day(g("дата"))
            if not day:
                continue
            fil = str(g("филиал") or "").strip()
            st = str(g("статус") or "").strip()
            # заготовка: менеджер протянул дату заранее, остальное пусто — это не лид, а пустая строка
            # (у Шынгыса 03.09 таких было 22 из 57 — раздували «строк в отчёте» и «нет телефона»)
            if not (norm_phone(g("телефон")) or str(g("фио") or "").strip() or fil or st):
                continue
            rows.append({"sheet": sn, "day": day, "name": str(g("фио") or "").strip()[:40],
                         "phone": norm_phone(g("телефон")), "filial": "" if fil.lower() == "none" else fil,
                         "status": st, "appt": str(g("дата записи") or "").strip(),
                         "comment": str(g("комментарий") or "").strip()[:80],
                         "prepay": str(g("предоплата") or "").strip()[:20],
                         "booked": bool(fil and fil.lower() != "none") or ("запис" in st.lower())})
    wb.close()
    data = {"taken": datetime.datetime.fromtimestamp(mt).strftime("%Y-%m-%d %H:%M"), "rows": rows, "sheets": sheets}
    _MEM.update(mtime=mt, data=data)
    return data
